#!/usr/bin/env python3
"""Generate the weekly EDA Jira ticket categorization workbook.

The script consumes the JSON returned by Jira's issue search API and updates a
copy of ``EDA_Ticket_Categorization.xlsx``.  Existing template category
assignments are retained; new tickets are categorized by deterministic rules.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from copy import copy
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import PatternFill


STATUS_ORDER = ("Open", "In Progress", "Resolved", "Closed")
CATEGORIES = (
    "Snowflake Provisioning",
    "AWS Provisioning",
    "CI/CD Setup / Github",
    "Analytics/Feature Request",
)
STATUS_COLORS = {
    "Open": "D9EAF7",
    "In Progress": "FFF2CC",
    "Resolved": "E2F0D9",
    "Closed": "D9D9D9",
}
RAW_SHEET = "Ticket Categorization"
SUMMARY_SHEET = "Summary"
RAW_HEADERS = ("Ticket", "Title", "Category", "Status")
SUMMARY_HEADERS = ("Category", "Count", *STATUS_ORDER)

CI_CD_TERMS = (
    "ci/cd",
    "github",
    "git hub",
    "repository",
    " repo",
)
AWS_TERMS = ("aws", "s3", "acre folder")
SNOWFLAKE_TERMS = (
    "snowflake",
    "warehouse",
    "schema",
    " role",
    "rbac",
    "polaris",
    "cortex",
    "external stage",
    "datashare",
    "data share",
    "posit",
    "ontology",
    "private listing",
    "iceberg",
)
DESCRIPTION_CATEGORY_PATTERN = re.compile(
    r"(?im)^\*\*Category:\*\*\s*([^\r\n]+)"
)


def value_name(value: Any) -> str:
    """Return a display value from a Jira field or normalized scalar."""
    if value is None:
        return ""
    if isinstance(value, dict):
        return str(
            value.get("name")
            or value.get("value")
            or value.get("displayName")
            or ""
        )
    return str(value)


def joined_names(value: Any) -> str:
    """Return comma-separated names from Jira list fields."""
    if not value:
        return ""
    if isinstance(value, str):
        return value
    return ", ".join(value_name(item) for item in value if value_name(item))


def description_category(value: Any) -> str:
    """Return a report category from Jira's structured description field."""
    if not isinstance(value, str):
        return ""
    match = DESCRIPTION_CATEGORY_PATTERN.search(value)
    if not match:
        return ""

    category = match.group(1).casefold()
    if "ci/cd" in category or "github" in category:
        return "CI/CD Setup / Github"
    if "aws" in category:
        return "AWS Provisioning"
    if "snowflake" in category:
        return "Snowflake Provisioning"
    return ""


def normalize_status(jira_status: str, status_category_key: str) -> str:
    """Map Jira statuses into the four report status groups."""
    if jira_status.casefold() == "closed":
        return "Closed"
    if status_category_key == "done":
        return "Resolved"
    if status_category_key == "indeterminate":
        return "In Progress"
    return "Open"


def classify_issue(fields: dict[str, Any], template_category: str = "") -> str:
    """Classify an issue into the category taxonomy used by the template."""
    if template_category:
        return template_category

    jira_category = value_name(fields.get("customfield_19840"))
    if jira_category:
        for category in CATEGORIES:
            if category.casefold() in jira_category.casefold():
                return jira_category

    structured_category = description_category(fields.get("description"))
    if structured_category:
        return structured_category

    text = " ".join(
        (
            str(fields.get("summary") or ""),
            joined_names(fields.get("components")),
            joined_names(fields.get("labels")),
        )
    ).casefold()
    if any(term in text for term in CI_CD_TERMS):
        return "CI/CD Setup / Github"
    if any(term in text for term in AWS_TERMS):
        return "AWS Provisioning"
    if any(term in text for term in SNOWFLAKE_TERMS):
        return "Snowflake Provisioning"
    return "Analytics/Feature Request"


def normalize_issue(
    issue: dict[str, Any], template_categories: dict[str, str]
) -> dict[str, str]:
    """Normalize either a Jira API issue or a pre-normalized issue object."""
    fields = issue.get("fields", issue)
    status = fields.get("status") or {}
    jira_status = value_name(status) or str(fields.get("jira_status") or "")
    category_key = (
        status.get("statusCategory", {}).get("key", "")
        if isinstance(status, dict)
        else ""
    ) or str(fields.get("status_category_key") or "")
    key = str(issue.get("key") or fields.get("key") or "")

    return {
        "key": key,
        "title": str(fields.get("summary") or fields.get("title") or ""),
        "status_group": str(
            fields.get("normalized_status")
            or normalize_status(jira_status, category_key)
        ),
        "category": classify_issue(fields, template_categories.get(key, "")),
    }


def load_issues(
    path: Path, template_categories: dict[str, str]
) -> list[dict[str, str]]:
    """Load and normalize issues from JSON."""
    with path.open(encoding="utf-8") as source:
        payload = json.load(source)
    issues = payload.get("issues", []) if isinstance(payload, dict) else payload
    normalized = [
        normalize_issue(issue, template_categories) for issue in issues
    ]
    if not normalized:
        raise ValueError("The input contains no Jira issues")
    if any(not issue["key"] for issue in normalized):
        raise ValueError("Every Jira issue must have a key")
    return sorted(normalized, key=lambda issue: issue["key"])


def template_category_assignments(ws: Any) -> dict[str, str]:
    """Return the manually curated ticket categories embedded in the template."""
    return {
        str(ticket): str(category)
        for ticket, _, category, _ in ws.iter_rows(
            min_row=2, min_col=1, max_col=4, values_only=True
        )
        if ticket and category
    }


def build_raw_sheet(ws: Any, issues: list[dict[str, str]]) -> None:
    """Replace the template's ticket rows while preserving its presentation."""
    header_styles = [copy(cell._style) for cell in ws[1][:4]]
    data_styles = [copy(cell._style) for cell in ws[2][:4]]
    data_height = ws.row_dimensions[2].height
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for column, header in enumerate(RAW_HEADERS, start=1):
        cell = ws.cell(1, column, header)
        cell._style = header_styles[column - 1]

    for issue in issues:
        ws.append(
            [
                issue["key"],
                issue["title"],
                issue["category"],
                issue["status_group"],
            ]
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row_number, row in enumerate(
        ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=4), start=2
    ):
        ws.row_dimensions[row_number].height = data_height
        for index, cell in enumerate(row):
            cell._style = copy(data_styles[index])
        key_cell, _, _, status_cell = row
        key_cell.hyperlink = f"https://wbdprod.atlassian.net/browse/{key_cell.value}"
        key_cell.style = "Hyperlink"
        status_cell.fill = PatternFill(
            "solid", fgColor=STATUS_COLORS[status_cell.value]
        )


def rollup_category(category: str) -> str:
    """Map template subcategories to one of the four summary categories."""
    for summary_category in CATEGORIES:
        if summary_category.casefold() in category.casefold():
            return summary_category
    return "Analytics/Feature Request"


def build_summary_sheet(ws: Any, issues: list[dict[str, str]]) -> None:
    """Build category-by-status counts using the template's summary layout."""
    header_styles = [copy(cell._style) for cell in ws[1][:4]]
    data_styles = [copy(cell._style) for cell in ws[2][:4]]
    header_height = ws.row_dimensions[1].height
    data_height = ws.row_dimensions[2].height
    ws.delete_rows(1, ws.max_row)

    for column, header in enumerate(SUMMARY_HEADERS, start=1):
        cell = ws.cell(1, column, header)
        cell._style = copy(
            header_styles[min(column - 1, len(header_styles) - 1)]
        )
    ws.row_dimensions[1].height = header_height

    counts = {
        category: Counter(
            issue["status_group"]
            for issue in issues
            if rollup_category(issue["category"]) == category
        )
        for category in CATEGORIES
    }
    for row_number, category in enumerate(CATEGORIES, start=2):
        category_counts = counts[category]
        values = [
            category,
            sum(category_counts.values()),
            *(category_counts[status] for status in STATUS_ORDER),
        ]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row_number, column, value)
            cell._style = copy(
                data_styles[min(column - 1, len(data_styles) - 1)]
            )
        ws.row_dimensions[row_number].height = data_height

    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 14


def generate_report(
    input_path: Path,
    output_path: Path,
    report_date: date,
    template_path: Path,
) -> None:
    wb = load_workbook(template_path)
    missing_sheets = {RAW_SHEET, SUMMARY_SHEET} - set(wb.sheetnames)
    if missing_sheets:
        raise ValueError(
            f"Template is missing required sheets: {sorted(missing_sheets)}"
        )

    template_categories = template_category_assignments(wb[RAW_SHEET])
    issues = load_issues(input_path, template_categories)
    build_raw_sheet(wb[RAW_SHEET], issues)
    build_summary_sheet(wb[SUMMARY_SHEET], issues)
    wb.properties.title = "EDA Ticket Categorization"
    wb.properties.subject = f"Weekly Jira report for {report_date.isoformat()}"
    wb.properties.creator = "EDA Jira report automation"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Jira issue JSON input")
    parser.add_argument("output", type=Path, help="Output .xlsx path")
    parser.add_argument(
        "--report-date",
        type=date.fromisoformat,
        default=date.today(),
        help="Report week date in YYYY-MM-DD format (default: today)",
    )
    parser.add_argument(
        "--template",
        type=Path,
        default=Path("EDA_Ticket_Categorization.xlsx"),
        help="Template workbook path",
    )
    return parser.parse_args()


if __name__ == "__main__":
    arguments = parse_args()
    generate_report(
        arguments.input,
        arguments.output,
        arguments.report_date,
        arguments.template,
    )
